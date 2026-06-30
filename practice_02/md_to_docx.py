#!/usr/bin/env python3
"""将古诗文汇编 Markdown 转为 Word 文档。"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).parent


def _set_default_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    style.paragraph_format.space_after = Pt(6)


def _set_heading_font(paragraph, size: int, bold: bool = True) -> None:
    for run in paragraph.runs:
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        run.font.size = Pt(size)
        run.font.bold = bold


def _add_runs_with_bold(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part:
            paragraph.add_run(part)


def md_to_docx(md_path: Path, docx_path: Path | None = None) -> Path:
    docx_path = docx_path or md_path.with_suffix(".docx")
    doc = Document()
    _set_default_font(doc)

    for level in (1, 2, 3):
        h = doc.styles[f"Heading {level}"]
        h.font.name = "黑体"
        h._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        h.font.bold = True
        h.font.size = Pt(18 - (level - 1) * 2)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]

        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        if line.strip().startswith("<a "):
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
            _set_heading_font(p, 18)
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=2)
            _set_heading_font(p, 16)
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=3)
            _set_heading_font(p, 14)
            i += 1
            continue

        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(12)
            _add_runs_with_bold(p, line[2:].strip())
            i += 1
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_bold(p, line[2:].strip())
            i += 1
            continue

        para_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if (
                not nxt.strip()
                or nxt.startswith("#")
                or nxt.startswith("> ")
                or nxt.startswith("- ")
                or nxt.strip() == "---"
                or nxt.strip().startswith("<a ")
            ):
                break
            para_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        _add_runs_with_bold(p, "\n".join(para_lines))

    doc.save(docx_path)
    return docx_path


def main() -> None:
    if len(sys.argv) > 1:
        md = Path(sys.argv[1])
    else:
        md = ROOT / "9A_古诗文汇编.md"
    out = md.with_suffix(".docx") if len(sys.argv) <= 2 else Path(sys.argv[2])
    path = md_to_docx(md, out)
    print(f"Wrote {path}")


if __name__ == "__main__":
    main()
